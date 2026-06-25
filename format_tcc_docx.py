"""Format TCC_final.docx without changing document text.

The script normalizes DOCX styles and layout, then writes a cleaned copy under
output/doc/. It intentionally avoids editing paragraph/run text.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


SOURCE = Path("TCC_final.docx")
OUTPUT = Path("output/doc/TCC_final_formatado.docx")


PAGE_BREAK_HEADINGS = {
    "AGRADECIMENTOS",
    "RESUMO",
    "ABSTRACT",
    "LISTA DE FIGURAS",
    "LISTA DE TABELAS",
    "LISTA DE ABREVIATURAS E SIGLAS",
    "SUMÁRIO",
    "1  INTRODUÇÃO",
    "2  O QUE A LITERATURA DIZ",
    "3  O QUE FIZEMOS",
    "4  RESULTADOS",
    "5  DISCUSSÃO",
    "6  LIMITAÇÕES",
    "7  CONCLUSÃO",
    "REFERÊNCIAS",
    "APÊNDICE A: VARIÁVEIS DO MODELO PRINCIPAL",
    "APÊNDICE B: DETALHES DE MODELAGEM",
    "APÊNDICE C: REPRODUTIBILIDADE",
}


def set_run_font(run, name="Times New Roman", size=12):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def set_style_font(style, name="Times New Roman", size=12, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def ensure_paragraph_style(styles, name, base=None):
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if base is not None:
            style.base_style = base
    return style


def clear_page_break_before(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    for node in list(ppr):
        if node.tag == qn("w:pageBreakBefore"):
            ppr.remove(node)


def remove_empty_nonimage_paragraphs(doc):
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip():
            continue
        if paragraph._p.xpath(".//w:drawing"):
            continue
        paragraph._p.getparent().remove(paragraph._p)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_styles(doc):
    styles = doc.styles
    normal = ensure_paragraph_style(styles, "Normal")
    set_style_font(normal, size=12, bold=False)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for name in ("Heading 1", "Title"):
        if name in styles:
            style = styles[name]
            set_style_font(style, size=14, bold=True)
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(12)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    heading2 = ensure_paragraph_style(styles, "Heading 2", normal)
    set_style_font(heading2, size=13, bold=True)
    heading2.paragraph_format.first_line_indent = Cm(0)
    heading2.paragraph_format.space_before = Pt(10)
    heading2.paragraph_format.space_after = Pt(6)
    heading2.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    for name in ("Legenda Figura", "Legenda Tabela"):
        caption = ensure_paragraph_style(styles, name, normal)
        set_style_font(caption, size=10, bold=False)
        caption.paragraph_format.first_line_indent = Cm(0)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(6)
        caption.paragraph_format.space_after = Pt(8)
        caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    return normal


def normalize_paragraphs(doc, normal):
    first_text_index = {}

    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        clear_page_break_before(paragraph)

        if not text:
            if paragraph._p.xpath(".//w:drawing"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = True
            continue

        first_text_index.setdefault(text, idx)
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name not in {"Heading 1", "Heading 2", "Legenda Figura", "Legenda Tabela"}:
            paragraph.style = normal

        is_title_page = idx <= 6
        is_presentation_page = 7 <= idx <= 14
        is_approval_page = 15 <= idx <= 25
        is_front_page = is_title_page or is_presentation_page or is_approval_page
        is_list_heading = text in {
            "LISTA DE FIGURAS",
            "LISTA DE TABELAS",
            "LISTA DE ABREVIATURAS E SIGLAS",
            "SUMÁRIO",
        }
        is_placeholder = text.startswith("Abra este documento no Microsoft Word")
        is_abbreviation = 40 <= idx <= 45
        is_caption = paragraph.style.name in {"Legenda Figura", "Legenda Tabela"} or text.startswith(
            ("Figura ", "Tabela ", "Fonte:")
        )

        paragraph.paragraph_format.page_break_before = (
            (text in PAGE_BREAK_HEADINGS and idx == first_text_index[text])
            or idx in {7, 15}
        )

        if is_front_page or is_list_heading:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            for run in paragraph.runs:
                set_run_font(run, size=12)
                run.font.bold = False
            if idx in {0, 1, 7, 15, 19} or is_list_heading:
                for run in paragraph.runs:
                    run.font.bold = True
            if idx in {2, 3, 8, 9, 16}:
                for run in paragraph.runs:
                    run.font.size = Pt(14)
                    run.font.bold = True
            if idx in {11, 12, 17}:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Cm(6.5)
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            if idx in {20, 22, 24}:
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.space_after = Pt(2)
            continue

        if paragraph.style.name == "Heading 1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=14)
                run.font.bold = True
            continue

        if paragraph.style.name == "Heading 2":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=13)
                run.font.bold = True
            continue

        if is_caption:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                set_run_font(run, size=10)
                run.font.bold = False
            if paragraph.runs:
                paragraph.runs[0].font.bold = True
            continue

        if is_placeholder:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            for run in paragraph.runs:
                set_run_font(run, size=10)
                run.font.bold = False
                run.font.italic = True
            continue

        if is_abbreviation:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

        for run in paragraph.runs:
            set_run_font(run, size=12)
            run.font.bold = False

        if text.startswith(("Palavras-chave:", "Keywords:")) and paragraph.runs:
            paragraph.runs[0].font.bold = True


def normalize_images(doc):
    max_width = Cm(15.5)
    for shape in doc.inline_shapes:
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = int(shape.height * ratio)


def normalize_tables(doc):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        try:
            table.style = "Table Grid"
        except KeyError:
            pass
        table.autofit = True
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                set_cell_borders(cell)
                set_cell_margins(cell)
                if row_idx == 0:
                    set_cell_shading(cell, "EDEDED")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.left_indent = Cm(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if len(paragraph.text.strip()) < 24
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        set_run_font(run, size=10)
                        run.font.bold = row_idx == 0


def main():
    doc = Document(SOURCE)

    for section in doc.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.start_type = WD_SECTION.NEW_PAGE

    normal = configure_styles(doc)
    remove_empty_nonimage_paragraphs(doc)
    normalize_paragraphs(doc, normal)
    normalize_images(doc)
    normalize_tables(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
